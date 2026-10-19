#!/usr/bin/env python3
"""Impacto da operação de nicotina sem fumaça em Santa Cruz do Sul (2037).

Calcula a cadeia com a inversa de Leontief da MIP-RS 2019 (DEE/SPGG e UFRJ),
nível 117 produtos × 52 atividades. A planilha oficial fica em /tmp e não entra
no repositório. Os resultados em R$ de 2025 saem em estudo/resultados/.

Três gavetas, nesta ordem:
  1. Efeito direto da planta (VAB, ocupações e remunerações do arquivo de premissas).
  2. Cadeia: só a compra feita no RS, a preço básico, choque nos produtos fornecedores.
  3. Vazamento para fora do RS, registrado e não multiplicado.

A receita ex-fábrica não é o choque. A atividade "fabricação de produtos do fumo"
não recebe a receita. Obra e CAPEX ficam fora.
"""

from __future__ import annotations

import json
import os
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import openpyxl
from matplotlib.patches import FancyBboxPatch

ROOT = Path(__file__).resolve().parents[1]
RESULTADOS = ROOT / "resultados"
FIGURAS = ROOT / "figuras"
TMP = Path("/tmp/mip_rs")
MIP_PATH = TMP / "mip_117_52.xlsx"
TRU_PATH = TMP / "tru_rs_2019.xlsx"
MIP_URL = (
    "https://dee.rs.gov.br/upload/arquivos/202409/"
    "11100149-20240906162755mip-rs-2019-117-52-v2-1.xlsx"
)
TRU_URL = (
    "https://dee.rs.gov.br/upload/arquivos/202409/"
    "11095935-20240906160547tru-rs-2019-v3-1.xlsx"
)

# IBGE SIDRA tabela 1737, variável 2266 (número-índice, dez/1993 = 100).
# Conferido em 19/10/2026: dezembro de 2019 e dezembro de 2025.
IPCA_DEZ_2019 = 5320.25
IPCA_DEZ_2025 = 7403.29
PHI = IPCA_DEZ_2019 / IPCA_DEZ_2025  # multiplica fluxo de 2025 para entrar na matriz

# Premissas da revisão 4 (18/10/2026), ano cheio 2037, R$ de 2025.
VBP_INT = 12.56e9
VBP_EXP = 3.27e9
VBP = VBP_INT + VBP_EXP  # 15,83 bi
BOM_INT = 5.81e9
BOM_EXP = 2.19e9
BOM = BOM_INT + BOM_EXP  # 8,00 bi
VA_DIRETO = 7.83e9  # 15,83 − 8,00
LEAF_TOTAL = 114.2e6
LEAF_RS = 52.8e6
LEAF_TON_TOTAL = 6510.0
LEAF_TON_RS = 3007.0
OCC_DIRETO = 4060.0
OCC_EXP = 1000.0
OCC_INT = OCC_DIRETO - OCC_EXP  # 3.060
REM_POR_OCUP = 150_000.0
REM_DIRETO = OCC_DIRETO * REM_POR_OCUP  # R$ 609 milhões
TAXA_CONTEUDO_LOCAL = (0.10, 0.20, 0.35)
TAXA_BASE = 0.20
PROPENSAO_CONSUMO_RS = 0.60

# Reconstrução das fichas v6 para separar interno e exportação.
# Volumes em unidades físicas; BOM em R$ por unidade; folha em t e R$/t.
ELIQ_INT_ML = 2.118e9
ELIQ_EXP_ML = 0.956e9
STICK_INT = 11.42e9
STICK_EXP = 0.918e9
SACH_EXP = 0.138e9
BOM_ELIQ = 2.14375
BOM_STICK = 0.110925
BOM_SACH = 0.268225
LEAF_G_ELIQ = 0.90
LEAF_G_STICK = 0.30
LEAF_G_SACH = 0.27
PRECO_LEAF_NIC = 15_000.0
PRECO_LEAF_REC = 19_500.0

# Bloco II do contratante (Tipo I deles, 2037, RS). Comparação, não insumo.
BLOCO_II = {
    "vab_rs_bi": 8.85,
    "ocupacoes_rs": 12760.0,
    "tributos_consumo_rs_bi": 0.928,
    "vab_brasil_bi": 15.55,
    "ocupacoes_brasil": 67350.0,
}

# Exemplos já escritos nas premissas (ranking de imprensa, a conferir no SIDRA).
# Não são estimativa de área de influência.
EXEMPLOS_FOLHA = [
    {"municipio": "Canguçu", "uf": "RS", "pam_t": 21064, "folha_nova_t": 181, "lat": -31.395, "lon": -52.676},
    {"municipio": "Venâncio Aires", "uf": "RS", "pam_t": 17779, "folha_nova_t": 153, "lat": -29.606, "lon": -52.192},
    {"municipio": "São Lourenço do Sul", "uf": "RS", "pam_t": 16374, "folha_nova_t": 141, "lat": -31.365, "lon": -51.978},
    {"municipio": "Santa Cruz do Sul", "uf": "RS", "pam_t": 12610, "folha_nova_t": 108, "lat": -29.718, "lon": -52.426, "sede": True},
    {"municipio": "Candelária", "uf": "RS", "pam_t": 11462, "folha_nova_t": 98, "lat": -29.669, "lon": -52.789},
]

# 52 atividades da MIP → 16 grupos da TRU (mesma família DEE/UFRJ 2019).
# A conferência com o valor da produção da TRU valida o mapa.
GRUPOS_16 = [
    ("Agropecuária", ["AT01", "AT02", "AT03"]),
    ("Indústria extrativa", ["AT04", "AT05"]),
    ("Indústrias de transformação", [f"AT{i:02d}" for i in range(6, 33)]),
    ("Eletricidade, gás, água e resíduos", ["AT33", "AT34"]),
    ("Construção", ["AT35"]),
    ("Comércio", ["AT36"]),
    ("Transporte, armazenagem e correio", ["AT37", "AT38", "AT39", "AT40"]),
    ("Alojamento e alimentação", ["AT41", "AT42"]),
    ("Informação e comunicação", ["AT43", "AT44"]),
    ("Atividades financeiras e de seguros", ["AT45"]),
    ("Atividades imobiliárias", ["AT46"]),
    ("Atividades profissionais e administrativas", ["AT47", "AT48"]),
    ("Administração pública", ["AT49"]),
    ("Educação e saúde mercantis", ["AT50"]),
    ("Artes, cultura e outros serviços", ["AT51"]),
    ("Serviços domésticos", ["AT52"]),
]


def baixar(url: str, destino: Path) -> None:
    if destino.exists() and destino.stat().st_size > 10_000:
        return
    destino.parent.mkdir(parents=True, exist_ok=True)
    import urllib.request

    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(req, timeout=120) as resp:
        destino.write_bytes(resp.read())


def norm(valor) -> str:
    if valor is None:
        return ""
    texto = str(valor).replace("-\n", "").replace("-\r", "")
    return " ".join(texto.replace("\n", " ").replace("\r", " ").split())


def ler_bloco(ws, n_linhas: int, n_cols: int, linha0: int = 5, col0: int = 2):
    """Lê rótulos da coluna A/B e um bloco numérico."""
    codigos = []
    nomes = []
    dados = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i < linha0:
            continue
        if len(codigos) >= n_linhas:
            break
        codigo = norm(row[0]) if row else ""
        if not codigo:
            continue
        codigos.append(codigo)
        nomes.append(norm(row[1]) if len(row) > 1 else "")
        linha = []
        for j in range(n_cols):
            v = row[col0 + j] if len(row) > col0 + j else 0
            linha.append(0.0 if v is None else float(v))
        dados.append(linha)
    return codigos, nomes, np.array(dados, dtype=float)


def cabecalhos_colunas(ws, n_cols: int, linha_cod: int = 3, linha_nome: int = 4, col0: int = 2):
    linhas = list(ws.iter_rows(max_row=linha_nome + 1, values_only=True))
    cod = linhas[linha_cod]
    nome = linhas[linha_nome]
    codigos, nomes = [], []
    for j in range(n_cols):
        codigos.append(norm(cod[col0 + j]) if len(cod) > col0 + j else "")
        nomes.append(norm(nome[col0 + j]) if len(nome) > col0 + j else "")
    return codigos, nomes


def carregar_mip():
    baixar(MIP_URL, MIP_PATH)
    baixar(TRU_URL, TRU_PATH)
    wb = openpyxl.load_workbook(MIP_PATH, read_only=True, data_only=True)

    prod_cod, prod_nom, B = ler_bloco(wb["9"], 117, 52)
    # Tabela 12 é atividades × produtos.
    ati_cod, ati_nom, D = ler_bloco(wb["12"], 52, 117)
    ati_a, _, A = ler_bloco(wb["13"], 52, 52)
    ati_l, _, L_pub = ler_bloco(wb["14"], 52, 52)
    if ati_cod != ati_a or ati_a != ati_l:
        raise SystemExit("Ordem das atividades diverge entre as tabelas 12, 13 e 14.")

    # Demanda da produção gaúcha a preço básico (tabela 3) e usos a preço de consumidor (tabela 2).
    dem_cod, dem_nom, dem = ler_bloco(wb["3"], 117, 62)
    uso_cod, _, uso = ler_bloco(wb["2"], 117, 52)
    col_dem, nome_dem = cabecalhos_colunas(wb["3"], 62)
    wb.close()

    if dem_cod != prod_cod or uso_cod != prod_cod:
        raise SystemExit("Ordem dos produtos diverge entre as tabelas 2, 3 e 9.")

    I = np.eye(52)
    L = np.linalg.inv(I - A)
    erro_L = np.max(np.abs(L - L_pub))
    erro_ident = np.max(np.abs((I - A) @ L - I))
    A_db = D @ B
    erro_A = np.max(np.abs(A_db - A))
    soma_D = D.sum(axis=0)

    # Colunas da tabela 3, depois das 52 atividades: totais e demanda final.
    idx = {norm(c).lower(): i for i, c in enumerate(col_dem)}
    # Os cabeçalhos de exportação e de consumo vêm na linha de nomes quando o código está vazio.
    # Mapeia pelo nome normalizado da linha 4 quando o código não identifica.
    nomes_l = [norm(n).lower() for n in nome_dem]

    def achar(*chaves):
        for i, (c, n) in enumerate(zip(col_dem, nomes_l)):
            texto = f"{c} {n}".lower()
            if all(k in texto for k in chaves):
                return i
        raise SystemExit(f"Coluna não encontrada: {chaves}")

    c_ci = achar("total")  # primeiro "Total" é o total do consumo intermediário
    # O segundo total é o das exportações. Procura de forma explícita.
    totais = [i for i, c in enumerate(col_dem) if norm(c).lower() == "total"]
    if len(totais) < 1:
        raise SystemExit(f"Totais não identificados: {col_dem}")
    c_ci = totais[0]
    c_fam = achar("fam")
    c_gov = achar("admin")
    c_isfl = achar("isfl")
    c_fbcf = achar("capital")
    c_est = achar("estoque")
    c_df = None
    for i, c in enumerate(col_dem):
        if "demanda" in norm(c).lower() and "final" in norm(c).lower():
            c_df = i
    if c_df is None:
        # linha de código pode estar quebrada: "Demanda final"
        for i, n in enumerate(nomes_l):
            if "demanda" in n and "final" in n:
                c_df = i
                break
    c_xw = achar("resto do mundo")
    c_xuf = achar("demais unidades")

    ci_prod = dem[:, c_ci]
    if c_df is not None:
        q = ci_prod + dem[:, c_df]
    else:
        q = (
            dem[:, :52].sum(axis=1)
            + dem[:, c_xw]
            + dem[:, c_xuf]
            + dem[:, c_gov]
            + dem[:, c_isfl]
            + dem[:, c_fam]
            + dem[:, c_fbcf]
            + dem[:, c_est]
        )
    q_alt = (
        dem[:, :52].sum(axis=1)
        + dem[:, c_xw]
        + dem[:, c_xuf]
        + dem[:, c_gov]
        + dem[:, c_isfl]
        + dem[:, c_fam]
        + dem[:, c_fbcf]
        + dem[:, c_est]
    )

    x_ativ = D @ q
    ci_ativ = uso.sum(axis=0)  # preço de consumidor, R$ milhões de 2019
    va_ativ = x_ativ - ci_ativ
    with np.errstate(divide="ignore", invalid="ignore"):
        coef_va = np.where(x_ativ > 1.0, va_ativ / x_ativ, 0.0)

    consumo = dem[:, c_fam].astype(float)
    consumo_pos = np.clip(consumo, 0.0, None)
    massa_neg = float((-np.clip(consumo, None, 0)).sum())
    if consumo_pos.sum() <= 0:
        raise SystemExit("Vetor de consumo das famílias está vazio.")
    share_consumo = consumo_pos / consumo_pos.sum()

    tru = carregar_tru()
    mapa_grupo = {}
    for nome_g, membros in GRUPOS_16:
        for m in membros:
            mapa_grupo[m] = nome_g
    if set(mapa_grupo) != set(ati_cod):
        raise SystemExit("Mapa de 16 grupos não cobre as 52 atividades.")

    prod_grupo = {g[0]: 0.0 for g in GRUPOS_16}
    for codigo, x in zip(ati_cod, x_ativ):
        prod_grupo[mapa_grupo[codigo]] += float(x)

    coef_rem = {}
    coef_occ = {}  # ocupações por R$ milhão de produção de 2019
    coef_va_tru = {}
    # A ordem de GRUPOS_16 segue a ordem das 16 colunas da TRU.
    for i, (nome_g, _) in enumerate(GRUPOS_16):
        prod = float(tru["producao"][i])
        coef_rem[nome_g] = float(tru["remuneracoes"][i] / prod)
        coef_occ[nome_g] = float(tru["ocupacoes"][i] / prod)
        coef_va_tru[nome_g] = float(tru["va"][i] / prod)

    w_ativ = np.array([coef_rem[mapa_grupo[c]] for c in ati_cod])
    e_ativ = np.array([coef_occ[mapa_grupo[c]] for c in ati_cod])

    return {
        "prod_cod": prod_cod,
        "prod_nom": prod_nom,
        "ati_cod": ati_cod,
        "ati_nom": ati_nom,
        "A": A,
        "B": B,
        "D": D,
        "L": L,
        "q": q,
        "q_alt": q_alt,
        "x_ativ": x_ativ,
        "va_ativ": va_ativ,
        "coef_va": coef_va,
        "ci_ativ": ci_ativ,
        "share_consumo": share_consumo,
        "consumo_pos_mi": float(consumo_pos.sum()),
        "consumo_neg_mi": massa_neg,
        "w_ativ": w_ativ,
        "e_ativ": e_ativ,
        "mapa_grupo": mapa_grupo,
        "prod_grupo": prod_grupo,
        "tru": tru,
        "coef_rem": coef_rem,
        "coef_occ": coef_occ,
        "coef_va_tru": coef_va_tru,
        "soma_D_min": float(soma_D.min()),
        "soma_D_max": float(soma_D.max()),
        "erro_L": float(erro_L),
        "erro_ident": float(erro_ident),
        "erro_A": float(erro_A),
        "col_dem": col_dem,
        "nome_dem": nome_dem,
        "idx_consumo": {
            "ci": int(c_ci),
            "familia": int(c_fam),
            "demanda_final": None if c_df is None else int(c_df),
        },
    }


def carregar_tru():
    wb = openpyxl.load_workbook(TRU_PATH, read_only=True, data_only=True)
    ws = wb["VA_16"]
    linhas = list(ws.iter_rows(values_only=True))
    wb.close()
    cab = [norm(c) for c in linhas[3][1:17]]
    def linha(rotulo_ini):
        for row in linhas:
            if norm(row[0]).lower().startswith(rotulo_ini):
                return [float(row[j]) for j in range(1, 17)]
        raise SystemExit(f"Linha TRU não encontrada: {rotulo_ini}")

    return {
        "grupos": cab,
        "va": linha("valor adicionado"),
        "remuneracoes": linha("remunera"),
        "producao": linha("valor da produ"),
        "ocupacoes": linha("fator trabalho"),
    }


def toneladas_folha(ml_eliq, n_stick, n_sach) -> float:
    return (
        ml_eliq * LEAF_G_ELIQ / 1e6
        + n_stick * LEAF_G_STICK / 1e6
        + n_sach * LEAF_G_SACH / 1e6
    )


def valor_folha(ml_eliq, n_stick, n_sach) -> float:
    t_eliq = ml_eliq * LEAF_G_ELIQ / 1e6
    t_stick = n_stick * LEAF_G_STICK / 1e6
    t_sach = n_sach * LEAF_G_SACH / 1e6
    return t_eliq * PRECO_LEAF_NIC + t_stick * PRECO_LEAF_REC + t_sach * PRECO_LEAF_NIC


def montar_compras(alpha: float, com_exportacao: bool) -> dict:
    """Compras no RS em R$ de 2025. Folha é dada. O resto segue a hipótese alpha."""
    folha_int = valor_folha(ELIQ_INT_ML, STICK_INT, 0.0)
    folha_exp = valor_folha(ELIQ_EXP_ML, STICK_EXP, SACH_EXP)
    share_int = folha_int / (folha_int + folha_exp)
    ton_int = toneladas_folha(ELIQ_INT_ML, STICK_INT, 0.0)
    ton_exp = toneladas_folha(ELIQ_EXP_ML, STICK_EXP, SACH_EXP)
    share_ton_int = ton_int / (ton_int + ton_exp)

    if com_exportacao:
        bom = BOM
        folha_total = LEAF_TOTAL
        folha_rs = LEAF_RS
        eliq_bom = (ELIQ_INT_ML + ELIQ_EXP_ML) * BOM_ELIQ
        toneladas_rs = LEAF_TON_RS
    else:
        bom = BOM_INT
        folha_total = LEAF_TOTAL * share_int
        folha_rs = LEAF_RS * share_int
        eliq_bom = ELIQ_INT_ML * BOM_ELIQ
        toneladas_rs = LEAF_TON_RS * share_ton_int

    nao_folha = bom - folha_total
    dispositivo_nacional = 0.50 * eliq_bom
    rs_dispositivo = 0.10 * dispositivo_nacional
    alvo = alpha * nao_folha
    if rs_dispositivo > alvo:
        rs_dispositivo = alvo
        rs_quim = 0.0
        rs_papel = 0.0
        ajuste = "dispositivo limitado pelo teto de conteúdo local"
    else:
        resto = alvo - rs_dispositivo
        # 30 e 40 são pesos relativos da hipótese de composição, não alíquotas de nota.
        rs_quim = resto * 30.0 / 70.0
        rs_papel = resto * 40.0 / 70.0
        ajuste = "residual do bolo repartido 30/70 químicos e 40/70 papel"

    compras = {
        "PR004": folha_rs,
        "PR066": rs_dispositivo,
        "PR051": rs_quim,
        "PR036": rs_papel,
    }
    return {
        "com_exportacao": com_exportacao,
        "alpha": alpha,
        "bom": bom,
        "folha_total": folha_total,
        "folha_rs": folha_rs,
        "toneladas_rs": toneladas_rs,
        "nao_folha": nao_folha,
        "alvo_rs_nao_folha": alvo,
        "dispositivo_nacional": dispositivo_nacional,
        "rs_dispositivo": rs_dispositivo,
        "rs_quimicos": rs_quim,
        "rs_papel": rs_papel,
        "compras_rs": sum(compras.values()),
        "vazamento": bom - sum(compras.values()),
        "ajuste": ajuste,
        "share_folha_interna": share_int,
        "por_produto": compras,
    }


def direto(com_exportacao: bool) -> dict:
    if com_exportacao:
        return {
            "vbp": VBP,
            "va": VA_DIRETO,
            "ocupacoes": OCC_DIRETO,
            "remuneracoes": REM_DIRETO,
            "bom": BOM,
        }
    va = VBP_INT - BOM_INT
    occ = OCC_INT
    return {
        "vbp": VBP_INT,
        "va": va,
        "ocupacoes": occ,
        "remuneracoes": occ * REM_POR_OCUP,
        "bom": BOM_INT,
    }


def propagar(mip, compras_rs: dict) -> dict:
    """Choque de produto em R$ de 2025 → produção, VAB, salários e ocupações da cadeia."""
    idx = {c: i for i, c in enumerate(mip["prod_cod"])}
    g_2025 = np.zeros(117)
    for cod, valor in compras_rs.items():
        g_2025[idx[cod]] += valor
    g_2019_mi = g_2025 * PHI / 1e6  # R$ milhões de 2019
    f = mip["D"] @ g_2019_mi
    x = mip["L"] @ f
    va = mip["coef_va"] * x
    rem = mip["w_ativ"] * x
    occ = mip["e_ativ"] * x
    return {
        "g_2025": g_2025,
        "f_2019_mi": f,
        "x_2019_mi": x,
        "va_2019_mi": va,
        "rem_2019_mi": rem,
        "occ": occ,
        "producao_2025": float(x.sum() * 1e6 / PHI),
        "va_2025": float(va.sum() * 1e6 / PHI),
        "rem_2025": float(rem.sum() * 1e6 / PHI),
        "ocupacoes": float(occ.sum()),
        "choque_atividade_2025": f * 1e6 / PHI,
    }


def induzido(mip, remuneracao_rs: float) -> dict:
    """Uma rodada. 60% da massa salarial do RS vira consumo no perfil da MIP."""
    consumo = PROPENSAO_CONSUMO_RS * remuneracao_rs
    g_2025 = mip["share_consumo"] * consumo
    g_2019_mi = g_2025 * PHI / 1e6
    f = mip["D"] @ g_2019_mi
    x = mip["L"] @ f
    va = mip["coef_va"] * x
    rem = mip["w_ativ"] * x
    occ = mip["e_ativ"] * x
    va_2025 = float(va.sum() * 1e6 / PHI)
    rem_2025 = float(rem.sum() * 1e6 / PHI)
    # Série fechada, só para nota: realimenta o salário desta rodada.
    k = (PROPENSAO_CONSUMO_RS * rem_2025 / consumo) if consumo else 0.0
    fator = (1.0 / (1.0 - k)) if 0 <= k < 1 else None
    return {
        "consumo_2025": consumo,
        "x_2019_mi": x,
        "va_2019_mi": va,
        "producao_2025": float(x.sum() * 1e6 / PHI),
        "va_2025": va_2025,
        "rem_2025": rem_2025,
        "ocupacoes": float(occ.sum()),
        "k_realimentacao": k,
        "fator_serie_fechada": fator,
        "va_serie_fechada_2025": None if fator is None else va_2025 * fator,
    }


def fiscal(com_exportacao: bool, va: float, rem: float) -> dict:
    """Fiscal bruto da revisão 4. Exportação é imune: o tributo sobre consumo não muda."""
    eob = va - rem
    irpj = 0.34 * 0.50 * eob
    return {
        "tributos_consumo_brasil": 13.70e9,
        "tributos_consumo_federal": 6.58e9,
        "tributos_consumo_subnacional_brasil": 7.12e9,
        "tributos_consumo_rs": 0.928e9,
        "tributos_consumo_rs_estadual_municipal": 0.48e9,
        "irpj_csll": irpj,
        "incentivo": 0.0,
        "retorno_liquido": None,
        "nota_irpj": (
            "mesma regra da revisão 4: 34% sobre lucro fiscal = 50% do excedente "
            "(VAB − remunerações). No cenário sem exportação, a regra é reaplicada "
            "ao VAB e à remuneração internos."
            if not com_exportacao
            else "hipótese da revisão 4, R$ 1,23 bi no ano cheio com exportação."
        ),
    }


def cenario(mip, alpha: float, com_exportacao: bool) -> dict:
    d = direto(com_exportacao)
    compras = montar_compras(alpha, com_exportacao)
    cadeia = propagar(mip, compras["por_produto"])
    ind = induzido(mip, d["remuneracoes"] + cadeia["rem_2025"])
    fisc = fiscal(com_exportacao, d["va"], d["remuneracoes"])
    return {
        "alpha": alpha,
        "com_exportacao": com_exportacao,
        "direto": d,
        "compras": {k: v for k, v in compras.items() if k != "por_produto"},
        "por_produto_2025": compras["por_produto"],
        "cadeia": {
            "producao_2025": cadeia["producao_2025"],
            "va_2025": cadeia["va_2025"],
            "rem_2025": cadeia["rem_2025"],
            "ocupacoes": cadeia["ocupacoes"],
        },
        "induzido": {
            "consumo_2025": ind["consumo_2025"],
            "producao_2025": ind["producao_2025"],
            "va_2025": ind["va_2025"],
            "rem_2025": ind["rem_2025"],
            "ocupacoes": ind["ocupacoes"],
            "k_realimentacao": ind["k_realimentacao"],
            "fator_serie_fechada": ind["fator_serie_fechada"],
            "va_serie_fechada_2025": ind["va_serie_fechada_2025"],
        },
        "total_tipo_I": {
            "va_2025": d["va"] + cadeia["va_2025"],
            "rem_2025": d["remuneracoes"] + cadeia["rem_2025"],
            "ocupacoes": d["ocupacoes"] + cadeia["ocupacoes"],
            "producao_2025": d["vbp"] + cadeia["producao_2025"],
        },
        "total_tipo_II": {
            "va_2025": d["va"] + cadeia["va_2025"] + ind["va_2025"],
            "rem_2025": d["remuneracoes"] + cadeia["rem_2025"] + ind["rem_2025"],
            "ocupacoes": d["ocupacoes"] + cadeia["ocupacoes"] + ind["ocupacoes"],
            "producao_2025": d["vbp"] + cadeia["producao_2025"] + ind["producao_2025"],
        },
        "fiscal": fisc,
        "_x_cadeia": cadeia["x_2019_mi"],
        "_va_cadeia": cadeia["va_2019_mi"],
        "_x_ind": ind["x_2019_mi"],
        "_va_ind": ind["va_2019_mi"],
        "_f_cadeia": cadeia["choque_atividade_2025"],
    }


def mi(valor: float) -> int:
    return int(round(valor / 1e6))


def bi1(valor: float) -> float:
    return round(valor / 1e9, 1)


def bi2(valor: float) -> float:
    return round(valor / 1e9, 2)


def salvar_tabelas(mip, base, faixa, sem_exp):
    RESULTADOS.mkdir(parents=True, exist_ok=True)
    # Vetor de choque por produto e por atividade (R$ milhões de 2025).
    linhas = ["codigo,nome,compra_rs_mi_2025,hipotese"]
    rotulo = {
        "PR004": "dado das premissas (folha, porteira, margem zero)",
        "PR066": "hipótese de composição: dispositivo",
        "PR051": "hipótese de composição: químicos",
        "PR036": "hipótese de composição: embalagem e papel",
    }
    for cod, nome in zip(mip["prod_cod"], mip["prod_nom"]):
        valor = base["por_produto_2025"].get(cod, 0.0)
        if valor == 0:
            continue
        linhas.append(
            f"{cod},\"{nome}\",{valor/1e6:.3f},{rotulo.get(cod, '')}"
        )
    (RESULTADOS / "choque_produtos.csv").write_text("\n".join(linhas) + "\n", encoding="utf-8")

    linhas = ["codigo,atividade,choque_mi_2025,producao_cadeia_mi_2025,va_cadeia_mi_2025,coef_va"]
    for i, (cod, nome) in enumerate(zip(mip["ati_cod"], mip["ati_nom"])):
        linhas.append(
            f"{cod},\"{nome}\","
            f"{base['_f_cadeia'][i]/1e6:.3f},"
            f"{base['_x_cadeia'][i] * 1e6 / PHI / 1e6:.3f},"
            f"{base['_va_cadeia'][i] * 1e6 / PHI / 1e6:.3f},"
            f"{mip['coef_va'][i]:.4f}"
        )
    (RESULTADOS / "cadeia_por_atividade.csv").write_text("\n".join(linhas) + "\n", encoding="utf-8")

    def linha_cen(nome, c):
        return (
            f"{nome},{c['alpha']:.2f},{int(c['com_exportacao'])},"
            f"{mi(c['direto']['va'])},{mi(c['cadeia']['va_2025'])},{mi(c['induzido']['va_2025'])},"
            f"{mi(c['total_tipo_I']['va_2025'])},{mi(c['total_tipo_II']['va_2025'])},"
            f"{c['direto']['ocupacoes']:.0f},{c['cadeia']['ocupacoes']:.0f},{c['induzido']['ocupacoes']:.0f},"
            f"{c['total_tipo_I']['ocupacoes']:.0f},{c['total_tipo_II']['ocupacoes']:.0f},"
            f"{mi(c['direto']['remuneracoes'])},{mi(c['cadeia']['rem_2025'])},{mi(c['induzido']['rem_2025'])},"
            f"{mi(c['compras']['compras_rs'])},{mi(c['compras']['vazamento'])},"
            f"{mi(c['cadeia']['producao_2025'])},{mi(c['induzido']['producao_2025'])}"
        )

    cab = (
        "cenario,alpha,com_exportacao,"
        "va_direto_mi,va_cadeia_mi,va_induzido_mi,va_tipo_I_mi,va_tipo_II_mi,"
        "occ_direto,occ_cadeia,occ_induzido,occ_tipo_I,occ_tipo_II,"
        "rem_direto_mi,rem_cadeia_mi,rem_induzido_mi,"
        "compras_rs_mi,vazamento_mi,producao_cadeia_mi,producao_induzido_mi"
    )
    linhas = [cab, linha_cen("base_20_com_exportacao", base)]
    for a, c in faixa.items():
        linhas.append(linha_cen(f"local_{int(a*100)}_com_exportacao", c))
    linhas.append(linha_cen("base_20_sem_exportacao", sem_exp))
    (RESULTADOS / "resultados_cenarios.csv").write_text("\n".join(linhas) + "\n", encoding="utf-8")

    linhas = ["grupo_tru16,producao_mip_mi_2019,producao_tru_mi_2019,coef_remuneracao,coef_ocupacoes_por_mi"]
    tru_prod = dict(zip(mip["tru"]["grupos"], mip["tru"]["producao"]))
    # A ordem dos nomes pode diferir; casa pelo índice da lista GRUPOS_16, que segue a TRU.
    for i, (nome, _) in enumerate(GRUPOS_16):
        tru_nome = mip["tru"]["grupos"][i]
        linhas.append(
            f"\"{nome}\",{mip['prod_grupo'][nome]:.3f},{mip['tru']['producao'][i]:.3f},"
            f"{mip['coef_rem'][nome]:.6f},{mip['coef_occ'][nome]:.6f}"
        )
    (RESULTADOS / "coeficientes_tru16.csv").write_text("\n".join(linhas) + "\n", encoding="utf-8")


def estilo():
    plt.rcParams.update(
        {
            "font.family": "DejaVu Sans",
            "axes.edgecolor": "#D5D1C8",
            "axes.labelcolor": "#1B3A4B",
            "xtick.color": "#1B3A4B",
            "ytick.color": "#1B3A4B",
            "text.color": "#1B3A4B",
            "figure.facecolor": "white",
            "axes.facecolor": "white",
            "axes.titleweight": "medium",
            "axes.titlesize": 12,
            "axes.labelsize": 10,
        }
    )


NAVY = "#1B3A4B"
TEAL = "#1F7A6B"
GOLD = "#B0893E"
CLAY = "#A24B32"
SLATE = "#5C6B73"


def figura_gavetas(base):
    estilo()
    fig, ax = plt.subplots(figsize=(8.4, 4.6))
    nomes = ["Direto\n(planta)", "Cadeia\n(compras no RS)", "Induzido\n(uma rodada)", "Total\nno RS"]
    vals = [
        base["direto"]["va"] / 1e9,
        base["cadeia"]["va_2025"] / 1e9,
        base["induzido"]["va_2025"] / 1e9,
        base["total_tipo_II"]["va_2025"] / 1e9,
    ]
    cores = [NAVY, TEAL, GOLD, "#0E2A36"]
    barras = ax.bar(nomes, vals, color=cores, width=0.72)
    for b, v in zip(barras, vals):
        ax.text(
            b.get_x() + b.get_width() / 2,
            v + 0.15,
            f"{v:.2f}".replace(".", ","),
            ha="center",
            va="bottom",
            fontsize=11,
            color=NAVY,
        )
    ax.set_ylabel("R$ bilhões de 2025")
    ax.set_title("Valor adicionado no Rio Grande do Sul em 2037\nCenário base, conteúdo local de 20% nos insumos não agrícolas")
    ax.set_ylim(0, max(vals) * 1.18)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    fig.tight_layout()
    fig.savefig(FIGURAS / "fig1_valor_adicionado.png", dpi=160)
    plt.close()


def figura_setores(mip, base):
    estilo()
    va = base["_va_cadeia"] * 1e6 / PHI / 1e9
    ordem = np.argsort(va)[::-1][:8]
    nomes = []
    for i in ordem:
        nome = mip["ati_nom"][i]
        if len(nome) > 42:
            nome = nome[:40] + "…"
        nomes.append(nome)
    valores = va[ordem]
    fig, ax = plt.subplots(figsize=(8.6, 5.0))
    y = np.arange(len(nomes))[::-1]
    ax.barh(y, valores, color=TEAL, height=0.7)
    ax.set_yticks(y)
    ax.set_yticklabels(nomes, fontsize=9)
    ax.set_xlabel("Valor adicionado da cadeia, R$ bilhões de 2025")
    ax.set_title("Onde a cadeia gaúcha acrescenta valor em 2037\nOito atividades com maior valor adicionado indireto")
    for yi, v in zip(y, valores):
        ax.text(v + 0.005, yi, f"{v:.2f}".replace(".", ","), va="center", fontsize=8, color=NAVY)
    ax.set_xlim(0, max(valores) * 1.22)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    fig.tight_layout()
    fig.savefig(FIGURAS / "fig2_cadeia_por_atividade.png", dpi=160)
    plt.close()


def figura_faixa(faixa, sem_exp, base):
    estilo()
    fig, axes = plt.subplots(1, 2, figsize=(9.2, 4.4))
    alphas = [10, 20, 35]
    va_c = [faixa[a / 100]["cadeia"]["va_2025"] / 1e9 for a in alphas]
    va_t = [faixa[a / 100]["total_tipo_II"]["va_2025"] / 1e9 for a in alphas]
    ax = axes[0]
    x = np.arange(len(alphas))
    ax.bar(x - 0.18, va_c, width=0.36, color=TEAL, label="Cadeia")
    ax.bar(x + 0.18, va_t, width=0.36, color=NAVY, label="Total com induzido")
    ax.set_xticks(x)
    ax.set_xticklabels([f"{a}%" for a in alphas])
    ax.set_xlabel("Conteúdo local dos insumos não agrícolas")
    ax.set_ylabel("R$ bilhões de 2025")
    ax.set_title("Faixa 10 / 20 / 35")
    ax.legend(frameon=False, fontsize=8)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)

    ax = axes[1]
    rotulos = ["Com exportação", "Sem exportação"]
    va_dir = [base["direto"]["va"] / 1e9, sem_exp["direto"]["va"] / 1e9]
    va_cad = [base["cadeia"]["va_2025"] / 1e9, sem_exp["cadeia"]["va_2025"] / 1e9]
    va_ind = [base["induzido"]["va_2025"] / 1e9, sem_exp["induzido"]["va_2025"] / 1e9]
    ax.bar(rotulos, va_dir, color=NAVY, label="Direto")
    ax.bar(rotulos, va_cad, bottom=va_dir, color=TEAL, label="Cadeia")
    baixo = [a + b for a, b in zip(va_dir, va_cad)]
    ax.bar(rotulos, va_ind, bottom=baixo, color=GOLD, label="Induzido")
    ax.set_ylabel("R$ bilhões de 2025")
    ax.set_title("Variante sem exportação")
    ax.legend(frameon=False, fontsize=8)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    fig.suptitle("Valor adicionado no RS em 2037", fontsize=12, color=NAVY)
    fig.tight_layout()
    fig.savefig(FIGURAS / "fig3_faixa_e_exportacao.png", dpi=160)
    plt.close()


def figura_comparacao(base):
    estilo()
    fig, ax = plt.subplots(figsize=(8.2, 4.4))
    nomes = [
        "Este estudo\nTipo I\n(direto + cadeia)",
        "Bloco II\ndo contratante\nTipo I, RS",
        "Este estudo\nTipo II\n(+ induzido)",
    ]
    vals = [
        base["total_tipo_I"]["va_2025"] / 1e9,
        BLOCO_II["vab_rs_bi"],
        base["total_tipo_II"]["va_2025"] / 1e9,
    ]
    cores = [TEAL, SLATE, NAVY]
    barras = ax.bar(nomes, vals, color=cores, width=0.62)
    for b, v in zip(barras, vals):
        ax.text(
            b.get_x() + b.get_width() / 2,
            v + 0.12,
            f"{v:.2f}".replace(".", ","),
            ha="center",
            fontsize=11,
        )
    ax.set_ylabel("Valor adicionado no RS, R$ bilhões de 2025")
    ax.set_title("Ordem de grandeza em 2037\nO bloco II não entra na conta e não se soma a este estudo")
    ax.set_ylim(0, max(vals) * 1.2)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    fig.tight_layout()
    fig.savefig(FIGURAS / "fig4_comparacao_bloco_ii.png", dpi=160)
    plt.close()


def figura_ocupacoes(base):
    estilo()
    fig, ax = plt.subplots(figsize=(8.2, 4.4))
    nomes = ["Direto\n(planta, premissa v6)", "Cadeia\n(coeficiente TRU-16)", "Induzido\n(coeficiente TRU-16)"]
    vals = [
        base["direto"]["ocupacoes"] / 1000,
        base["cadeia"]["ocupacoes"] / 1000,
        base["induzido"]["ocupacoes"] / 1000,
    ]
    cores = [NAVY, TEAL, GOLD]
    barras = ax.bar(nomes, vals, color=cores, width=0.68)
    for b, v in zip(barras, vals):
        ax.text(
            b.get_x() + b.get_width() / 2,
            v + 0.08,
            f"{v:.1f}".replace(".", ","),
            ha="center",
            fontsize=11,
        )
    ax.set_ylabel("Mil ocupações em 2037")
    ax.set_title("Ocupações no Rio Grande do Sul em 2037\nCadeia e induzido usam o coeficiente agregado da TRU, não um levantamento da planta")
    ax.set_ylim(0, max(vals) * 1.22)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    fig.tight_layout()
    fig.savefig(FIGURAS / "fig5_ocupacoes.png", dpi=160)
    plt.close()


def figura_mapa():
    estilo()
    fig, ax = plt.subplots(figsize=(8.2, 6.2))
    geo_ok = False
    geo_path = Path("/tmp/brazil-states.geojson")
    if not geo_path.exists():
        try:
            import urllib.request

            url = "https://raw.githubusercontent.com/codeforamerica/click_that_hood/master/public/data/brazil-states.geojson"
            req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
            with urllib.request.urlopen(req, timeout=30) as resp:
                geo_path.write_bytes(resp.read())
        except Exception:
            geo_ok = False
    if geo_path.exists():
        try:
            geo = json.loads(geo_path.read_text(encoding="utf-8"))
            for feat in geo["features"]:
                nome = (
                    feat.get("properties", {}).get("name")
                    or feat.get("properties", {}).get("sigla")
                    or ""
                )
                if nome not in {"Rio Grande do Sul", "RS"}:
                    continue
                geom = feat["geometry"]
                polys = geom["coordinates"] if geom["type"] == "MultiPolygon" else [geom["coordinates"]]
                if geom["type"] == "Polygon":
                    polys = [geom["coordinates"]]
                for poly in polys:
                    anel = poly[0]
                    xs = [p[0] for p in anel]
                    ys = [p[1] for p in anel]
                    ax.fill(xs, ys, color="#E7F0EA", zorder=0)
                    ax.plot(xs, ys, color="#7FA392", lw=0.8, zorder=1)
                geo_ok = True
        except Exception:
            geo_ok = False
    if not geo_ok:
        ax.set_xlim(-58, -49)
        ax.set_ylim(-34, -27)
        ax.add_patch(
            FancyBboxPatch(
                (-57.2, -33.4),
                7.4,
                6.2,
                boxstyle="round,pad=0.02,rounding_size=0.4",
                facecolor="#E7F0EA",
                edgecolor="#7FA392",
            )
        )
    desloc = {
        "Santa Cruz do Sul": (8, -14),
        "Venâncio Aires": (8, 8),
        "Candelária": (-78, -4),
        "Canguçu": (-62, -12),
        "São Lourenço do Sul": (8, -10),
    }
    for p in EXEMPLOS_FOLHA:
        sede = p.get("sede", False)
        ax.scatter(
            [p["lon"]],
            [p["lat"]],
            s=90 if sede else 46,
            color=CLAY if sede else NAVY,
            zorder=3,
        )
        texto = "Santa Cruz do Sul (sede)" if sede else p["municipio"]
        ax.annotate(
            texto,
            (p["lon"], p["lat"]),
            textcoords="offset points",
            xytext=desloc.get(p["municipio"], (6, 6)),
            fontsize=8,
            color=NAVY,
        )
    ax.set_title(
        "Sede e exemplos de compra de folha citados nas premissas\n"
        "Ilustração dos pontos já escritos no arquivo. Não é área de influência estimada."
    )
    ax.set_xlabel("Longitude")
    ax.set_ylabel("Latitude")
    ax.set_aspect("equal", adjustable="datalim")
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    fig.tight_layout()
    fig.savefig(FIGURAS / "fig6_mapa_exemplos.png", dpi=160)
    plt.close()


def limpar(obj):
    if isinstance(obj, dict):
        return {k: limpar(v) for k, v in obj.items() if not k.startswith("_")}
    if isinstance(obj, list):
        return [limpar(v) for v in obj]
    if isinstance(obj, float):
        return round(obj, 6) if abs(obj) < 1e6 else round(obj, 2)
    return obj


def main():
    FIGURAS.mkdir(parents=True, exist_ok=True)
    RESULTADOS.mkdir(parents=True, exist_ok=True)
    print("Carregando MIP-RS 2019...")
    mip = carregar_mip()
    print(f"PHI IPCA = {PHI:.6f}")
    print(f"erro L vs tabela 14 = {mip['erro_L']:.3e}")
    print(f"erro (I-A)L - I = {mip['erro_ident']:.3e}")
    print(f"erro D@B - A = {mip['erro_A']:.3e}")
    print(f"soma colunas D: {mip['soma_D_min']:.4f} a {mip['soma_D_max']:.4f}")
    print(f"produção total atividades R$ mi 2019 = {mip['x_ativ'].sum():.1f}")
    print(f"TRU produção total = {sum(mip['tru']['producao']):.1f}")
    print("Produção por grupo (MIP agregado vs TRU):")
    for i, (nome, _) in enumerate(GRUPOS_16):
        tru_n = mip["tru"]["grupos"][i]
        a = mip["prod_grupo"][nome]
        b = mip["tru"]["producao"][i]
        print(f"  {a:12.1f} {b:12.1f} dif {a-b:10.1f}  mip={nome[:28]:28s} tru={tru_n[:36]}")

    base = cenario(mip, TAXA_BASE, True)
    faixa = {a: cenario(mip, a, True) for a in TAXA_CONTEUDO_LOCAL}
    sem_exp = cenario(mip, TAXA_BASE, False)

    print("\n=== BASE 2037 ===")
    for bloco in ("direto", "cadeia", "induzido"):
        b = base[bloco]
        if bloco == "direto":
            print(f"direto VA {b['va']/1e9:.3f} occ {b['ocupacoes']:.0f} rem {b['remuneracoes']/1e6:.1f}")
        else:
            print(
                f"{bloco} VA {b['va_2025']/1e9:.3f} prod {b['producao_2025']/1e9:.3f} "
                f"occ {b['ocupacoes']:.0f} rem {b['rem_2025']/1e6:.1f}"
            )
    print("tipo I", base["total_tipo_I"]["va_2025"] / 1e9, "tipo II", base["total_tipo_II"]["va_2025"] / 1e9)
    print("compras RS", base["compras"]["compras_rs"] / 1e6, "vazamento", base["compras"]["vazamento"] / 1e6)
    print("k induzido", base["induzido"]["k_realimentacao"], "fator", base["induzido"]["fator_serie_fechada"])
    print("sem exp VA tipo II", sem_exp["total_tipo_II"]["va_2025"] / 1e9, "occ", sem_exp["total_tipo_II"]["ocupacoes"])
    for a, c in faixa.items():
        print(
            f"alpha {a:.0%} cadeia VA {c['cadeia']['va_2025']/1e9:.3f} "
            f"total II {c['total_tipo_II']['va_2025']/1e9:.3f} occ II {c['total_tipo_II']['ocupacoes']:.0f}"
        )

    # Coeficientes de VA dos setores que recebem o choque direto de atividade.
    print("\nChoque por atividade (R$ mi 2025) — maiores:")
    ordem = np.argsort(base["_f_cadeia"])[::-1][:8]
    for i in ordem:
        print(f"  {mip['ati_cod'][i]} {mip['ati_nom'][i][:50]:50s} {base['_f_cadeia'][i]/1e6:8.1f} va_coef {mip['coef_va'][i]:.3f}")

    salvar_tabelas(mip, base, faixa, sem_exp)
    figura_gavetas(base)
    figura_setores(mip, base)
    figura_faixa(faixa, sem_exp, base)
    figura_comparacao(base)
    figura_ocupacoes(base)
    figura_mapa()

    saida = {
        "matriz": {
            "nome": "MIP-RS 2019",
            "fonte": "DEE/SPGG e UFRJ (MAEP)",
            "agregacao": "117 produtos x 52 atividades",
            "arquivo": MIP_URL,
            "unidade_matriz": "R$ milhões correntes de 2019",
            "leontief": "L = (I − A)^(−1) calculada da tabela 13 e conferida com a tabela 14",
            "erro_max_L_vs_tabela14": mip["erro_L"],
            "erro_max_identidade": mip["erro_ident"],
            "erro_max_DxB_vs_A": mip["erro_A"],
            "ipca_dez_2019": IPCA_DEZ_2019,
            "ipca_dez_2025": IPCA_DEZ_2025,
            "phi_2019_sobre_2025": PHI,
            "formula": "fluxo_2019 = fluxo_2025 × (IPCA_dez2019 / IPCA_dez2025); resultado_2025 = resultado_2019 / phi",
        },
        "hipoteses": {
            "conteudo_local_nao_agricola_base": TAXA_BASE,
            "faixa": list(TAXA_CONTEUDO_LOCAL),
            "propensao_consumo_rs": PROPENSAO_CONSUMO_RS,
            "adicionalidade": 1.0,
            "capex": "fora desta rodada",
            "emprego_indireto": "coeficiente de ocupações por R$ milhão de produção da TRU-RS 2019, 16 atividades, aplicado a cada atividade do grupo",
        },
        "base": limpar(base),
        "faixa": {str(a): limpar(c) for a, c in faixa.items()},
        "sem_exportacao": limpar(sem_exp),
        "bloco_II_contratante": BLOCO_II,
        "exemplos_folha_premissas": EXEMPLOS_FOLHA,
        "consumo_familias_rs_preco_basico_mi_2019": mip["consumo_pos_mi"],
        "consumo_familias_valores_negativos_cortados_mi_2019": mip["consumo_neg_mi"],
    }
    (RESULTADOS / "resultados.json").write_text(
        json.dumps(saida, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print("Arquivos em", RESULTADOS, "e", FIGURAS)


if __name__ == "__main__":
    main()
